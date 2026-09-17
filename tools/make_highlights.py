"""Write paper/highlights.docx from the highlights block of cas_calmnet.tex, so
the separate highlights file always matches the manuscript. Elsevier limits each
highlight to 85 characters including spaces; the script refuses longer ones.

    python tools/make_highlights.py
PDF: convert highlights.docx with Word (tools/make_highlights_pdf.ps1).
"""
import io
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

PAPER = Path(__file__).resolve().parent.parent / "paper"


def tex_to_text(s):
    s = " ".join(s.split())
    s = s.replace("{,}", ",").replace("--", "\u2013").replace("\\%", "%")
    return re.sub(r"[{}]", "", s)


def main():
    tex = io.open(PAPER / "cas_calmnet.tex", encoding="utf-8").read()
    block = tex[tex.index("\\begin{highlights}"):tex.index("\\end{highlights}")]
    items = [tex_to_text(x) for x in block.split("\\item")[1:]]
    title = tex_to_text(re.search(r"\\title\[mode = title\]\{(.*?)\}\s*\n\s*\n", tex, re.S).group(1))
    for it in items:
        assert len(it) <= 85, (len(it), it)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    doc.add_heading("Highlights", level=1)
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    for it in items:
        doc.add_paragraph(it, style="List Bullet")
    doc.save(str(PAPER / "highlights.docx"))
    for it in items:
        print("%2d  %s" % (len(it), it))


if __name__ == "__main__":
    main()
