"""Midterm presentation deck: standard academic wording, built on the proposal's
own slides as templates (shapes, fonts, sizes, colours and geometry copied).
Tables are crops of the compiled paper; figures are the paper's figure files.

    python presentation/crop_paper_assets.py
    python presentation/build_midterm_presentation.py
"""
import copy
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
RES = HERE.parent / "results"
ASSETS = HERE / "paper_assets"
SRC = HERE / "CALM-Net_midterm_proposal.pptx"
OUT = HERE / (sys.argv[1] if len(sys.argv) > 1 else "CALM-Net_midterm_presentation.pptx")

BLK = RGBColor(0x11, 0x11, 0x11)
GRY = RGBColor(0x55, 0x55, 0x55)
FONT = "Times New Roman"
BUL = "•  "

prs = Presentation(str(SRC))
N_ORIG = len(prs.slides)
T = {i + 1: s for i, s in enumerate(prs.slides)}


def clone(ti, skip=()):
    src = T[ti]
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in src.shapes:
        if shp.shape_type in (13, 19) or shp.name in skip:
            continue
        new.shapes._spTree.insert_element_before(copy.deepcopy(shp._element), "p:extLst")
    return new


def borrow(s, ti, name, left=None, top=None, width=None, height=None):
    src = next(sh for sh in T[ti].shapes if sh.name == name)
    s.shapes._spTree.insert_element_before(copy.deepcopy(src._element), "p:extLst")
    sh = s.shapes[-1]
    place(sh, left, top, width, height)
    return sh


def get(s, name):
    return next(sh for sh in s.shapes if sh.name == name)


def place(sh, left=None, top=None, width=None, height=None):
    for attr, v in (("left", left), ("top", top), ("width", width), ("height", height)):
        if v is not None:
            setattr(sh, attr, Inches(v))


def R(text, size, bold=False, italic=False, color=BLK):
    return (text, size, bold, italic, color)


def P(*runs, align=PP_ALIGN.LEFT, sa=8, ls=1.1):
    return dict(runs=list(runs), align=align, sa=sa, ls=ls)


def write(shape, paras):
    tf = shape.text_frame
    body = tf._txBody
    for p in body.findall(qn("a:p"))[1:]:
        body.remove(p)
    first = tf.paragraphs[0]
    for child in list(first._p):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            first._p.remove(child)
    for i, spec in enumerate(paras):
        p = first if i == 0 else tf.add_paragraph()
        p.alignment, p.space_after, p.line_spacing = spec["align"], Pt(spec["sa"]), spec["ls"]
        for text, size, bold, italic, color in spec["runs"]:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
            f.color.rgb = color


def title(s, text):
    write(get(s, "TextBox 2"), [P(R(text, 30, bold=True), sa=8, ls=1.0)])


def picture(s, path, left, top, width, height):
    w, h = Image.open(path).size
    scale = min(width / w, height / h)
    pw, ph = w * scale, h * scale
    return s.shapes.add_picture(str(path), Inches(left + (width - pw) / 2),
                                Inches(top + (height - ph) / 2), Inches(pw), Inches(ph))


def lead(bold, rest, size):
    return P(R(bold, size, bold=True), R(rest, size, color=GRY), sa=8, ls=1.2)


def bullets(s, items, size=19):
    write(get(s, "TextBox 4"), [P(R(BUL + b, size, bold=True), R(t, size, color=GRY), sa=13, ls=1.18)
                                for b, t in items])


def cards(s, items):
    for name, (num, head, body) in zip(("Rectangle 4", "Rectangle 5", "Rectangle 6"), items):
        write(get(s, name), [P(R(num, 30, bold=True), sa=6, ls=1.1), P(R(head, 20, bold=True), sa=10, ls=1.05),
                             P(R(body, 16, color=GRY), sa=8, ls=1.22)])


def figure_slide(heading, path, bold, rest, top=1.55, height=4.60):
    s = clone(6)
    title(s, heading)
    picture(s, path, 0.87, top, 11.60, height)
    write(get(s, "TextBox 5"), [P(R(bold, 16, bold=True), R(rest, 16, color=GRY), sa=8, ls=1.15)])
    return s


def table_slide(heading, subtitle, image, box, note=None, points=(), points_top=5.4):
    s = clone(9, skip=("Rectangle 6", "Rectangle 7", "Rectangle 8") + (() if note else ("TextBox 9",)))
    title(s, heading)
    write(get(s, "TextBox 4"), [P(R(subtitle, 13.5, italic=True, color=GRY))])
    picture(s, image, *box)
    if note:
        nb = get(s, "TextBox 9")
        place(nb, top=note[0])
        write(nb, [P(R(note[1], 11.5, italic=True, color=GRY))])
    tk = get(s, "TextBox 10")
    place(tk, top=points_top, height=7.2 - points_top)
    write(tk, [lead(b, t, 16) for b, t in points])
    return s


# ------------------------------------------------------------------ 1 title
s = clone(1)
write(get(s, "TextBox 4"), [P(R("Label-Free Covariance Alignment for Longitudinal EEG Decoding", 24),
                              align=PP_ALIGN.CENTER, sa=8, ls=1.15),
                            P(R("Midterm Presentation", 18, color=GRY), align=PP_ALIGN.CENTER, sa=8, ls=1.15)])
write(get(s, "TextBox 5"), [P(R("Md Wahiduzzaman Suva  [26-94088-2]", 15, color=GRY), align=PP_ALIGN.CENTER, sa=5),
                            P(R("Esme Moula Chowdhury Abha  [26-94089-2]", 15, color=GRY), align=PP_ALIGN.CENTER, sa=5),
                            P(R("Supervisor: Dr. Muhammad Hasibur Rashid Chayon", 15, color=GRY), align=PP_ALIGN.CENTER, sa=8)])

# ------------------------------------------------------------------ 2 introduction
s = clone(2)
title(s, "Introduction")
write(get(s, "TextBox 4"), [P(R("Lower-limb exoskeletons controlled by non-invasive EEG can restore walking for people "
                                "with spinal cord injury or stroke. A decoder is trained once but must operate across "
                                "sessions recorded over several weeks.", 20), sa=8, ls=1.25)])
write(get(s, "Rectangle 5"), [P(R("Changes in electrode placement and impedance between sessions alter the EEG signal "
                                  "statistics that a trained decoder relies on.", 22, bold=True), sa=8, ls=1.15)])
write(get(s, "TextBox 6"), [P(R("This work:  ", 19, bold=True),
                              R("a compact EEG decoder that adapts to each new session without labels, evaluated on "
                                "three EEG datasets.", 19, color=GRY), sa=8, ls=1.2)])

# ------------------------------------------------------------------ 3 research challenges
s = clone(3)
title(s, "Research Challenges")
cards(s, [("1", "Session drift", "Electrode position and impedance change between sessions, shifting the signal "
                                "statistics a trained decoder depends on."),
          ("2", "Movement artefact", "Walking produces motion and muscle activity that can mix with the EEG "
                                     "recording."),
          ("3", "Reliable decisions", "Exoskeleton control needs consistent accuracy across sessions and "
                                      "well-calibrated confidence.")])

# ------------------------------------------------------------------ 4 objectives
s = clone(4)
title(s, "Research Objectives")
for name, head, body in (
        ("TextBox 5", "Adaptive alignment", "Design an alignment layer that corrects session drift without labels, "
                                            "during training and at test time."),
        ("TextBox 7", "Accurate and calibrated decoding", "Achieve accuracy comparable to published decoders with a "
                                                          "compact model and calibrated confidence."),
        ("TextBox 9", "A design rule for adaptation", "Derive and validate a rule, read from the recording protocol, "
                                                      "that sets the adaptation rate of the alignment layer.")):
    write(get(s, name), [P(R(head, 21, bold=True), sa=2, ls=1.1), P(R(body, 16, color=GRY), sa=8, ls=1.12)])
write(get(s, "TextBox 10"), [P(R("Research question:  ", 17, bold=True),
                               R("can a compact decoder adapt to session drift without labels while remaining accurate "
                                 "and reliable?", 17, italic=True, color=GRY))])

# ------------------------------------------------------------------ 5 literature review
s = clone(5)
title(s, "Literature Review")
boxes = {
    "Rectangle 4": ("Exoskeleton BCI", ["NeuroRex dataset  [Sarkar 2026]", "Deep-learning control  [Ferrero 2024]",
                                        "Error potentials  [Soriano 2025]", "Robot-assisted gait  [Tortora 2023]"]),
    "Rectangle 5": ("EEG decoding models", ["EEGNet  [Lawhern 2018]", "ShallowFBCSPNet  [Schirrmeister 2017]",
                                            "EEG Conformer  [Song 2023]"]),
    "Rectangle 6": ("Domain alignment", ["Euclidean alignment  [He 2020]", "Riemannian Procrustes  [Rodrigues 2019]",
                                         "Riemannian classifiers  [Barachant 2012]"]),
    "Rectangle 7": ("Test-time adaptation", ["Entropy minimisation  [Wang 2021]", "Temporal correlation  [Gong 2022]",
                                             "Selective prediction  [Geifman 2019]"]),
}
for name, (head, items) in boxes.items():
    write(get(s, name), [P(R(head, 15, bold=True), sa=9, ls=1.05)] +
          [P(R(it, 12, color=GRY), sa=7, ls=1.1) for it in items])
write(get(s, "Rectangle 8"), [P(R("RESEARCH GAP", 14, bold=True), sa=4, ls=1.1),
                              P(R("Existing alignment methods operate offline on complete recordings. This work places "
                                  "alignment inside the network, adapts it online without labels, and provides a rule "
                                  "for choosing its adaptation rate.", 19), sa=8, ls=1.14)])
write(get(s, "TextBox 9"), [P(R("Citations shown as [first author, year]; full entries are listed on the References slide.",
                                10.5, italic=True, color=GRY))])

# ------------------------------------------------------------------ 6 architecture
figure_slide("Proposed Architecture", ASSETS / "fig_arch_paper.png",
             "Proposed decoder (24,181 parameters):  ",
             "an adaptive alignment layer, a multi-scale power stem, an optional cross-epoch context module and a "
             "selective classification head.")

# ------------------------------------------------------------------ 7 method
s = clone(7)
title(s, "Adaptive Alignment Method")
bullets(s, [("Running covariance estimate:", " an exponential moving average of the spatial covariance of incoming EEG windows."),
            ("Whitening:", " each window is multiplied by the inverse square root of that estimate before feature extraction."),
            ("Label-free adaptation:", " the estimate continues to update on new sessions without labels."),
            ("Adaptation rate rule:", " the estimator memory is set longer than the class block length of the recording protocol.")])
write(get(s, "Rectangle 5"), [P(R("KEY IDEA   ", 14, bold=True),
                                R("Session drift is corrected inside the network, and the recording protocol determines the "
                                  "adaptation rate before training.", 18, italic=True), sa=8, ls=1.15)])

# ------------------------------------------------------------------ 8 datasets
s = clone(3)
title(s, "Datasets and Experimental Setup")
cards(s, [("A", "Exoskeleton EEG  [Sarkar 2026]", "7 participants, 9 sessions over several weeks. Trained on sessions "
                                                  "1 to 3 and tested on sessions 4 to 9."),
          ("B", "Treadmill walking  [Luu 2017]", "8 participants with 3 trials each, recorded in an independent "
                                                 "laboratory."),
          ("C", "Motor execution  [Schalk 2004]", "20 participants from PhysioNet EEGMMIDB, with the analysis plan "
                                                  "registered before training.")])
tb = borrow(s, 2, "TextBox 6", left=0.70, top=5.85, width=12.0, height=1.1)
write(tb, [P(R("Protocol:  ", 17, bold=True),
             R("8 to 30 Hz filtering, AdamW optimisation with early stopping, and three data-split seeds; seven "
               "published decoders trained with identical settings.", 17, color=GRY), sa=8, ls=1.2)])

# ------------------------------------------------------------------ 9 drift
figure_slide("Session Drift Reduction", RES / "fig_drift.png", "Drift reduced in all 15 participants:  ",
             "the alignment layer reduces the covariance distance between sessions by 45.9 % on cohort A and 62.5 % on "
             "cohort B, while a control without updates leaves it unchanged.")

# ------------------------------------------------------------------ 10 per-participant drift
s = clone(10, skip=("Rectangle 6", "Rectangle 7", "Rectangle 8", "Rectangle 9"))
title(s, "Per-Participant Drift Reduction")
write(get(s, "TextBox 4"), [P(R("Riemannian distance between the training-session covariance and the test sessions, before "
                                "and after alignment.", 13.5, italic=True, color=GRY))])
picture(s, ASSETS / "tab_drift.png", 0.70, 1.95, 11.93, 3.95)
write(get(s, "TextBox 10"), [lead("Cohort A:  ", "38.7 to 51.2 % reduction for each participant.", 15),
                             lead("Cohort B:  ", "43.8 to 82.1 % reduction for each participant.", 15)])

# ------------------------------------------------------------------ 11 comparison
table_slide("Comparison with Baseline Models",
            "Balanced accuracy (Acc) and expected calibration error (ECE), mean over three data-split seeds; all models "
            "trained in the same pipeline.",
            ASSETS / "tab_compare_ppt.png", (0.70, 1.95, 11.93, 3.35),
            note=(5.30, "† below 0.55 accuracy for at least five of twenty participants on one seed."),
            points=[("Cohort A:  ", "0.868 balanced accuracy, comparable to the best baseline (0.876), with fewer "
                                    "parameters than six of the seven baselines."),
                    ("Cohorts C and B:  ", "highest accuracy on cohort C (0.796); on cohort B, with alignment set by "
                                           "the protocol rule, 0.729, second of eight models.")],
            points_top=5.62)

# ------------------------------------------------------------------ 12 alignment estimate
figure_slide("Alignment Estimate During Recording", RES / "fig_class_tracking.png",
             "Estimate trajectory, one participant per cohort:  ",
             "with short class blocks (cohort A) the estimate stays between the class covariances; with long class "
             "blocks (cohort B) it follows the active class, the case the adaptation rate rule addresses.",
             top=1.45, height=4.85)

# ------------------------------------------------------------------ 13 adaptation rate
table_slide("Adaptation Rate Analysis",
            "Cohort B, three data-split seeds: balanced accuracy of align + gate at four adaptation rates.",
            ASSETS / "tab_ratecurve_ppt.png", (3.4, 2.00, 6.5, 2.85),
            points=[("Effect of the rate:  ", "lengthening the estimator memory from 160 to 3200 windows raises accuracy "
                                              "from 0.580 to 0.698, as the adaptation rate rule predicts."),
                    ("Protocol rule:  ", "when class blocks are longer than the estimator memory, the rule selects the "
                                         "configuration without alignment, which reaches 0.729 on this cohort.")],
            points_top=5.10)

# ------------------------------------------------------------------ 14 prospective validation
table_slide("Prospective Validation on Cohort C",
            "Predictions registered before any model was trained on cohort C; balanced accuracy, seed 0.",
            ASSETS / "tab_prereg.png", (1.2, 2.10, 10.9, 1.90),
            points=[("Both predictions confirmed:  ", "enabling alignment preserved accuracy (+0.002), and a slower "
                                                      "adaptation rate brought no change, as the rule predicts for short "
                                                      "class blocks."),
                    ("Replication:  ", "two additional data-split seeds support the same conclusions.")],
            points_top=4.45)

# ------------------------------------------------------------------ 15 ablation
table_slide("Ablation Study",
            "Cohort A, three data-split seeds.",
            ASSETS / "tab_ablation3_ppt.png", (1.8, 1.95, 9.7, 2.60),
            note=(4.65, "FA/min: spurious activations per minute of standing. Acc@90: accuracy at 90 % coverage."),
            points=[("Alignment layer:  ", "adding alignment raises balanced accuracy from 0.842 (gate only) to 0.868."),
                    ("Cross-epoch context:  ", "reduces spurious activations from 1.60 to 1.16 per minute of standing.")],
            points_top=5.10)

# ------------------------------------------------------------------ 16 artefact analysis
s = clone(8)
title(s, "Artefact Analysis")
picture(s, ASSETS / "tab_artefact.png", 0.55, 2.05, 6.30, 4.20)
write(get(s, "TextBox 5"), [
    P(R("Accuracy after ICA artefact removal", 18, bold=True), sa=12, ls=1.15),
    P(R("Eye, muscle and other non-neural components were removed with ICLabel [Pion-Tonachini 2019] before "
        "retraining the decoder.", 17, color=GRY), sa=12, ls=1.18),
    P(R("Mean balanced accuracy remains at ", 17, color=GRY), R("0.785", 17, bold=True),
      R(", well above chance, and the change for each participant is unrelated to the number of muscle components "
        "removed.", 17, color=GRY), sa=12, ls=1.18),
    P(R("The results are consistent with decoding from neural EEG activity.", 18, bold=True), sa=8, ls=1.18)])

# ------------------------------------------------------------------ 17 topography
figure_slide("Scalp Topography", RES / "fig_topography.png", "One participant, 8 to 30 Hz power:  ",
             "power decreases over sensorimotor electrodes during walking, consistent with mu and beta "
             "desynchronisation during movement.", top=1.60, height=4.50)

# ------------------------------------------------------------------ 18 limitations
s = clone(7)
title(s, "Limitations and Future Work")
bullets(s, [("Cohort size:", " the three datasets include 7, 8 and 20 participants."),
            ("Offline evaluation:", " all results are from recorded datasets; closed-loop testing is planned."),
            ("Architecture development:", " the architecture is under active development, and further variants will be evaluated."),
            ("Adaptation rate selection:", " the rate is currently set by a rule and could be learned from data.")])
write(get(s, "Rectangle 5"), [P(R("FUTURE WORK   ", 14, bold=True),
                                R("Real-time evaluation with an exoskeleton, larger cohorts, and recording protocols "
                                  "designed to refine the adaptation rate rule.", 18, italic=True), sa=8, ls=1.15)])

# ------------------------------------------------------------------ 19 conclusion
s = clone(2)
title(s, "Conclusion")
write(get(s, "TextBox 4"), [P(R("This work presents a compact EEG decoder with an adaptive alignment layer that corrects "
                                "session drift without labels.", 20), sa=8, ls=1.25)])
write(get(s, "Rectangle 5"), [P(R("The decoder reduces session drift by 46 % to 63 %, achieves accuracy comparable to "
                                  "published decoders, and follows a protocol-based rule confirmed prospectively.",
                                  22, bold=True), sa=8, ls=1.15)])
write(get(s, "TextBox 6"), [P(R("Outlook:  ", 19, bold=True),
                              R("the approach supports longitudinal use of EEG-controlled exoskeletons with less "
                                "recalibration.", 19, color=GRY), sa=8, ls=1.2)])

# ------------------------------------------------------------------ 20 references
s = clone(11)
write(get(s, "TextBox 4"), [P(R("Works cited on these slides, keyed to their [author, year] tags.", 12.5, italic=True, color=GRY))])
left = [
    ("[Sarkar 2026]", "S. Sarkar et al., “EEG-controlled exoskeleton for walking and standing: a longitudinal multimodal dataset,” Sci. Data, 2026. (OpenNeuro ds007788)"),
    ("[Ferrero 2024]", "L. Ferrero et al., “Brain-machine interface based on deep learning to control a lower-limb robotic exoskeleton,” J. NeuroEng. Rehabil., 21(48), 2024."),
    ("[Soriano 2025]", "P. Soriano-Segura, M. Ortiz et al., “Characterization of error-related potentials during exoskeleton command via deep learning,” J. NeuroEng. Rehabil., 2025."),
    ("[Tortora 2023]", "S. Tortora et al., “Cortical and muscular activity under robot-assisted gait modes during exoskeleton walking,” 2023."),
    ("[Luu 2017]", "T. P. Luu et al., “Real-time EEG-based brain-computer interface to a virtual avatar enhances cortical involvement in human treadmill walking,” Sci. Rep., 7:8895, 2017."),
    ("[Schalk 2004]", "G. Schalk et al., “BCI2000: a general-purpose brain-computer interface (BCI) system,” IEEE TBME, 51(6):1034-1043, 2004."),
    ("[Lawhern 2018]", "V. J. Lawhern et al., “EEGNet: a compact CNN for EEG-based brain-computer interfaces,” J. Neural Eng., 15(5), 2018."),
    ("[Schirrmeister 2017]", "R. T. Schirrmeister et al., “Deep learning with convolutional neural networks for EEG decoding and visualization,” Hum. Brain Mapp., 38(11):5391-5420, 2017."),
]
right = [
    ("[Song 2023]", "Y. Song et al., “EEG Conformer: convolutional transformer for EEG decoding and visualization,” IEEE TNSRE, 31:710-719, 2023."),
    ("[He 2020]", "H. He and D. Wu, “Transfer learning for brain-computer interfaces: a Euclidean space data alignment approach,” IEEE TBME, 67(2):399-410, 2020."),
    ("[Rodrigues 2019]", "P. L. C. Rodrigues et al., “Riemannian Procrustes analysis: transfer learning for brain-computer interfaces,” IEEE TBME, 66(8):2390-2401, 2019."),
    ("[Barachant 2012]", "A. Barachant et al., “Multiclass brain-computer interface classification by Riemannian geometry,” IEEE TBME, 59(4), 2012."),
    ("[Wang 2021]", "D. Wang et al., “Tent: fully test-time adaptation by entropy minimization,” ICLR, 2021."),
    ("[Gong 2022]", "T. Gong et al., “NOTE: robust continual test-time adaptation against temporal correlation,” NeurIPS, 2022."),
    ("[Geifman 2019]", "Y. Geifman and R. El-Yaniv, “SelectiveNet: a deep neural network with an integrated reject option,” ICML, 2019."),
    ("[Pion-Tonachini 2019]", "L. Pion-Tonachini et al., “ICLabel: an automated electroencephalographic independent component classifier, dataset, and website,” NeuroImage, 198:181-197, 2019."),
]
for name, refs in (("TextBox 5", left), ("TextBox 6", right)):
    write(get(s, name), [P(R(tag + "  ", 11, bold=True), R(txt, 10.5, color=GRY), sa=9, ls=1.06) for tag, txt in refs])

# ------------------------------------------------------------------ drop the original proposal slides
ids = prs.slides._sldIdLst
for sid in list(ids)[:N_ORIG]:
    prs.part.drop_rel(sid.rId)
    ids.remove(sid)
prs.save(str(OUT))
print("saved %s | %d slides" % (OUT.name, len(prs.slides)))
