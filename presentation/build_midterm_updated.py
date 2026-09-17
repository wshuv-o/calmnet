"""Update the CALM-Net midterm proposal deck with the results.

The proposal's own slides are the templates: each new slide is a copy of one of
them (shapes, fonts, sizes, colours, geometry), with only the text and images
changed. Tables are screenshots cropped from the compiled paper
(crop_paper_assets.py), the architecture is the paper's Figure 1, and the other
figures are the paper's figure files.

    python presentation/crop_paper_assets.py
    python presentation/build_midterm_updated.py
"""
import copy
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
RES = HERE.parent / "results"
ASSETS = HERE / "paper_assets"
SRC = HERE / "CALM-Net_midterm_proposal.pptx"
OUT = HERE / "CALM-Net_midterm_updated.pptx"

BLK = RGBColor(0x11, 0x11, 0x11)
GRY = RGBColor(0x55, 0x55, 0x55)
FONT = "Times New Roman"
MINUS = "−"

prs = Presentation(str(SRC))
N_ORIG = len(prs.slides)
T = {i + 1: s for i, s in enumerate(prs.slides)}


# ------------------------------------------------------------------ helpers
def clone(ti, skip=()):
    """New slide copying template slide ti, without pictures and tables."""
    src = T[ti]
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in src.shapes:
        if shp.shape_type in (13, 19) or shp.name in skip:
            continue
        new.shapes._spTree.insert_element_before(copy.deepcopy(shp._element), "p:extLst")
    return new


def borrow(s, ti, name, left=None, top=None, width=None, height=None):
    """Copy one shape from template slide ti into slide s."""
    src = next(sh for sh in T[ti].shapes if sh.name == name)
    el = copy.deepcopy(src._element)
    s.shapes._spTree.insert_element_before(el, "p:extLst")
    sh = s.shapes[-1]
    for attr, v in (("left", left), ("top", top), ("width", width), ("height", height)):
        if v is not None:
            setattr(sh, attr, Inches(v))
    return sh


def get(s, name):
    return next(sh for sh in s.shapes if sh.name == name)


def place(sh, left=None, top=None, width=None, height=None):
    for attr, v in (("left", left), ("top", top), ("width", width), ("height", height)):
        if v is not None:
            setattr(sh, attr, Inches(v))


def R(text, size, bold=False, italic=False, color=BLK):
    return (text, size, bold, italic, color)


def P(*runs, align=PP_ALIGN.LEFT, sa=8, ls=1.1):
    return dict(runs=list(runs), align=align, sa=sa, ls=ls)


def write(shape, paras):
    """Replace a shape's text, keeping the shape and its frame settings."""
    tf = shape.text_frame
    body = tf._txBody
    for p in body.findall(qn("a:p"))[1:]:
        body.remove(p)
    first = tf.paragraphs[0]
    for child in list(first._p):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            first._p.remove(child)
    for i, spec in enumerate(paras):
        p = first if i == 0 else tf.add_paragraph()
        p.alignment = spec["align"]
        p.space_after = Pt(spec["sa"])
        p.line_spacing = spec["ls"]
        for text, size, bold, italic, color in spec["runs"]:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
            f.color.rgb = color


def title(s, text):
    write(get(s, "TextBox 2"), [P(R(text, 30, bold=True), sa=8, ls=1.0)])


def picture(s, path, left, top, width, height):
    """Add an image fitted inside the box, centred, aspect kept."""
    w, h = Image.open(path).size
    scale = min(width / w, height / h)
    pw, ph = w * scale, h * scale
    return s.shapes.add_picture(str(path), Inches(left + (width - pw) / 2),
                                Inches(top + (height - ph) / 2), Inches(pw), Inches(ph))


def lead(bold, rest, size):
    return P(R(bold, size, bold=True), R(rest, size, color=GRY), sa=8, ls=1.2)


# ------------------------------------------------------------------ 1 title
s = clone(1)
write(get(s, "TextBox 4"), [P(R("Label-free covariance alignment for longitudinal EEG decoding", 24),
                              align=PP_ALIGN.CENTER, sa=8, ls=1.15),
                            P(R("Midterm proposal, updated with results", 18, color=GRY),
                              align=PP_ALIGN.CENTER, sa=8, ls=1.15)])
write(get(s, "TextBox 5"), [P(R("Md Wahiduzzaman Suva  [26-94088-2]", 15, color=GRY), align=PP_ALIGN.CENTER, sa=5),
                            P(R("Esme Moula Chowdhury Abha  [26-94089-2]", 15, color=GRY), align=PP_ALIGN.CENTER, sa=5),
                            P(R("Supervisor: Dr. Muhammad Hasibur Rashid Chayon", 15, color=GRY), align=PP_ALIGN.CENTER, sa=8)])

# ------------------------------------------------------------------ 2, 3 unchanged
clone(2)
clone(3)

# ------------------------------------------------------------------ 4 objectives and outcomes
s = clone(4)
title(s, "Objectives and Outcomes")
for name, head, body in (
        ("TextBox 5", "Movement-invariant intent",
         "Outcome: the planned leakage score was itself confounded. A conditional probe and an ICA "
         "retraining control show that most of the accuracy is neural."),
        ("TextBox 7", "Calibrated across time",
         "Outcome: the main contribution. A label-free alignment layer adapts at test time and removes "
         "46 % and 63 % of session drift."),
        ("TextBox 9", "Abstain with a guarantee",
         "Outcome: reported as negative. The trained selective head adds no accuracy over three seeds and "
         "lowers accuracy at 90 % coverage in 14 of 18 arms.")):
    write(get(s, name), [P(R(head, 21, bold=True), sa=2, ls=1.1), P(R(body, 16, color=GRY), sa=8, ls=1.12)])
write(get(s, "TextBox 10"), [P(R("Question answered:  ", 17, bold=True),
                               R("when does test-time alignment help a longitudinal EEG decoder, and when does it harm it?",
                                 17, italic=True, color=GRY))])

# ------------------------------------------------------------------ 5 literature review
s = clone(5)
boxes = {
    "Rectangle 4": ("Lower-limb exoskeleton BCI", ["NeuroRex dataset  [Sarkar 2026]", "Deep-learning control  [Ferrero 2024]",
                                                   "Exo error potentials  [Soriano 2025]", "Robot-assisted gait  [Tortora 2023]"]),
    "Rectangle 5": ("EEG decoders", ["EEGNet  [Lawhern 2018]", "ShallowFBCSPNet  [Schirrmeister 2017]",
                                     "EEG-Conformer  [Song 2023]"]),
    "Rectangle 6": ("Alignment and transfer", ["Euclidean alignment  [He 2020]", "Riemannian Procrustes  [Rodrigues 2019]",
                                               "Riemannian geom.  [Barachant 2012]"]),
    "Rectangle 7": ("Test-time adaptation", ["Tent  [Wang 2021]", "Correlated streams  [Gong 2022]",
                                             "Reject option  [Geifman 2019]"]),
}
for name, (head, items) in boxes.items():
    write(get(s, name), [P(R(head, 15, bold=True), sa=9, ls=1.05)] +
          [P(R(it, 12, color=GRY), sa=7, ls=1.1) for it in items])
write(get(s, "Rectangle 8"), [P(R("RESEARCH GAP", 14, bold=True), sa=4, ls=1.1),
                              P(R("Alignment corrects session drift offline, and test-time adaptation degrades on "
                                  "temporally correlated streams. No prior work states when a label-free alignment "
                                  "layer inside an EEG decoder helps or harms under a blocked recording protocol.", 19),
                                sa=8, ls=1.14)])
write(get(s, "TextBox 9"), [P(R("Citations shown as [first author, year]; full entries are on the References slide.",
                                10.5, italic=True, color=GRY))])

# ------------------------------------------------------------------ 6 architecture
s = clone(6)
title(s, "Decoder Architecture")
picture(s, ASSETS / "fig_arch_paper.png", 0.87, 1.55, 11.60, 4.60)
write(get(s, "TextBox 5"), [P(R("Proposed configuration, align + gate, 24,181 parameters:  ", 16, bold=True),
                              R("a covariance alignment layer that keeps updating on unlabelled test data, a multi-scale "
                                "power stem and a selective head. The cross-epoch context is optional and adds 594,816 "
                                "parameters.", 16, color=GRY), sa=8, ls=1.15)])

# ------------------------------------------------------------------ 7 core idea
s = clone(7)
title(s, "Core Idea and Hypothesis")
bul = [("•  Align inside the network:", " a running spatial covariance whitens every window and keeps updating on unlabelled test data."),
       ("•  Estimator memory:", " the momentum m sets a memory of about 32/m windows, 160 at the default m = 0.2."),
       ("•  The failure mode:", " when one class streams for longer than that memory, the estimate becomes that class and whitens the difference away."),
       ("•  A rule before training:", " the memory must outlast a class block, which is read from the recording protocol.")]
write(get(s, "TextBox 4"), [P(R(b, 19, bold=True), R(t, 19, color=GRY), sa=13, ls=1.18) for b, t in bul])
write(get(s, "Rectangle 5"), [P(R("HYPOTHESIS   ", 14, bold=True),
                                R("Alignment helps where class blocks are shorter than the estimator's memory and harms "
                                  "where they are longer, and the protocol tells which case applies before training.",
                                  18, italic=True), sa=8, ls=1.15)])

# ------------------------------------------------------------------ 8 three cohorts
s = clone(3)
title(s, "Three Cohorts, One Pipeline")
for name, num, head, body in (
        ("Rectangle 4", "A", "Exoskeleton  [Sarkar 2026]",
         "7 participants, 9 sessions over weeks. Fitted on sessions 1 to 3, tested on 4 to 9. Class blocks of 18 windows."),
        ("Rectangle 5", "B", "Treadmill  [Luu 2017]",
         "8 participants, 3 trials, an independent laboratory. Class blocks of 309 windows, longer than the estimator's memory."),
        ("Rectangle 6", "C", "Motor execution  [Schalk 2004]",
         "20 participants of PhysioNet EEGMMIDB. Class blocks of 5 windows. Predictions committed before any model was trained.")):
    write(get(s, name), [P(R(num, 30, bold=True), sa=6, ls=1.1), P(R(head, 20, bold=True), sa=10, ls=1.05),
                         P(R(body, 16, color=GRY), sa=8, ls=1.22)])
tb = borrow(s, 2, "TextBox 6", left=0.70, top=5.85, width=12.0, height=1.1)
write(tb, [P(R("Same settings for every model:  ", 17, bold=True),
             R("8 to 30 Hz, AdamW, early stopping, temperature scaling and three data-split seeds, with eight "
               "published decoders trained in the identical pipeline.", 17, color=GRY), sa=8, ls=1.2)])

# ------------------------------------------------------------------ 9 what is decoded
s = clone(8)
title(s, "What Is Actually Decoded")
picture(s, ASSETS / "tab_artefact.png", 0.55, 2.05, 6.30, 4.20)
write(get(s, "TextBox 5"), [
    P(R("The midterm confound score was itself confounded.", 18, bold=True), sa=12, ls=1.15),
    P(R("A decoder with no movement input scored ", 17, color=GRY), R("R² = 0.863", 17, bold=True),
      R(", because walk and stop differ in movement by definition.", 17, color=GRY), sa=12, ls=1.18),
    P(R("Retraining after ICA artefact removal [Pion-Tonachini 2019]: ", 17, color=GRY), R("0.835 to 0.785", 17, bold=True),
      R(" (p = 0.22). The losses do not follow the muscle components removed.", 17, color=GRY), sa=12, ls=1.18),
    P(R("Most of the accuracy is neural, and artefact is not fully excluded.", 18, bold=True), sa=8, ls=1.18)])

# ------------------------------------------------------------------ 10 topography
s = clone(6)
title(s, "Scalp Topography")
picture(s, RES / "fig_topography.png", 0.87, 1.60, 11.60, 4.50)
write(get(s, "TextBox 5"), [P(R("One participant, 8 to 30 Hz power:  ", 16, bold=True),
                              R("power falls over sensorimotor electrodes during walking, as mu and beta desynchronisation "
                                "predicts, and rises over occipital electrodes. The two effects are spatially separate.",
                                16, color=GRY), sa=8, ls=1.15)])

# ------------------------------------------------------------------ 11 drift
s = clone(6)
title(s, "Session Drift Removed")
picture(s, RES / "fig_drift.png", 0.87, 1.55, 11.60, 4.60)
write(get(s, "TextBox 5"), [P(R("15 of 15 participants:  ", 16, bold=True),
                              R("alignment reduces session-to-session covariance distance by 45.9 % on cohort A and "
                                "62.5 % on cohort B, while the no-op control leaves it unchanged.", 16, color=GRY),
                              sa=8, ls=1.15)])

# ------------------------------------------------------------------ 12 per-participant drift
s = clone(10, skip=("Rectangle 6", "Rectangle 7", "Rectangle 8", "Rectangle 9"))
title(s, "Per-Participant Drift")
write(get(s, "TextBox 4"), [P(R("Riemannian distance between the fitting-session covariance and the held-out sessions, "
                                "before and after alignment.", 13.5, italic=True, color=GRY))])
picture(s, ASSETS / "tab_drift.png", 0.70, 1.95, 11.93, 3.95)
write(get(s, "TextBox 10"), [lead("Cohort A:  ", "a reduction of 38.7 to 51.2 % per participant, and no participant regresses.", 15),
                             lead("Cohort B:  ", "43.8 to 82.1 %. The no-op control stays within 0.1 % of the raw distance "
                                                 "for every participant.", 15)])

# ------------------------------------------------------------------ 13 comparison
s = clone(9, skip=("Rectangle 6", "Rectangle 7", "Rectangle 8"))
title(s, "Comparison with Published Decoders")
write(get(s, "TextBox 4"), [P(R("Balanced accuracy (Acc) and calibration error (ECE), mean over three data-split seeds; "
                                "all nine decoders trained in one pipeline.", 13.5, italic=True, color=GRY))])
picture(s, ASSETS / "tab_compare.png", 0.70, 1.95, 11.93, 3.35)
note = get(s, "TextBox 9")
place(note, top=5.30)
write(note, [P(R("† near chance for at least five of twenty participants, a training failure under the shared settings.",
                 11.5, italic=True, color=GRY))])
tk = get(s, "TextBox 10")
place(tk, top=5.65, height=1.5)
write(tk, [lead("Takeaway:  ", "level with the published decoders on cohort A (0.868 against 0.817 to 0.878), highest on "
                              "cohort C (0.796), and level on cohort B once alignment is off (0.729 against 0.617 to 0.734).", 16),
           lead("No size advantage:  ", "EEG-TCNet is smaller. The contribution is the alignment layer and the rule "
                                         "that switches it.", 16)])

# ------------------------------------------------------------------ 14 class tracking
s = clone(6)
title(s, "The Estimator Tracks the Streaming Class")
picture(s, RES / "fig_class_tracking.png", 0.87, 1.45, 11.60, 4.85)
write(get(s, "TextBox 5"), [P(R("One participant per cohort:  ", 16, bold=True),
                              R("on cohort A the estimate stays between the classes; on cohort B it moves onto the class "
                                "being streamed and follows it through the walk block. Slowing it slows the migration.",
                                16, color=GRY), sa=8, ls=1.15)])

# ------------------------------------------------------------------ 15 rate across cohorts
s = clone(6)
title(s, "Adaptation Rate Across Cohorts")
picture(s, RES / "fig_rate.png", 0.87, 1.55, 11.60, 4.65)
write(get(s, "TextBox 5"), [P(R("(c) Effect of alignment against memory, three seeds:  ", 16, bold=True),
                              R("cohort B's cost shrinks as the memory passes its block length and never turns positive; "
                                "cohorts A and C stay near zero.", 16, color=GRY), sa=8, ls=1.15)])

# ------------------------------------------------------------------ 16 cohort B cost
s = clone(9, skip=("Rectangle 6", "Rectangle 7", "Rectangle 8", "TextBox 9"))
title(s, "Cost of Alignment on Cohort B")
write(get(s, "TextBox 4"), [P(R("Cohort B, three data-split seeds: align + gate minus gate only (0.729), paired within "
                                "participant.", 13.5, italic=True, color=GRY))])
picture(s, ASSETS / "tab_ratecurve.png", 2.2, 2.05, 8.9, 2.70)
tk = get(s, "TextBox 10")
place(tk, top=5.05, height=1.8)
write(tk, [lead("Takeaway:  ", "alignment costs 0.149 at the default rate, lower in 8 of 8 participants. Slowing "
                              "adaptation recovers 0.118, with no step at the 309-window block length.", 16),
           lead("At slow rates:  ", "the remaining cost of about 0.03 is not significant, and the layer never becomes "
                                     "beneficial on this cohort.", 16)])

# ------------------------------------------------------------------ 17 registered prediction
s = clone(9, skip=("Rectangle 6", "Rectangle 7", "Rectangle 8", "TextBox 9"))
title(s, "A Prediction Registered in Advance")
write(get(s, "TextBox 4"), [P(R("Cohort C, seed 0, exactly as committed to version control before any model was trained "
                                "on it.", 13.5, italic=True, color=GRY))])
picture(s, ASSETS / "tab_prereg.png", 1.2, 2.10, 10.9, 1.90)
tk = get(s, "TextBox 10")
place(tk, top=4.45, height=2.3)
write(tk, [lead("Both predictions held:  ", "enabling alignment cost nothing (+0.002) and slowing it brought no gain "
                                           "(" + MINUS + "0.006), where slowing had recovered 0.090 on cohort B.", 16),
           lead("Two further seeds agree:  ", "alignment +0.014 (p = 0.17); slowing costs 0.016, lower in 15 of 20 "
                                               "participants (p = 0.003).", 16)])

# ------------------------------------------------------------------ 18 slow adaptation
s = clone(7)
title(s, "Slow Adaptation Has a Cost")
bul = [("•  Cohort A:", " slowing from m = 0.2 to 0.01 lowers accuracy from 0.868 to 0.819, lower in 7 of 7 participants (p = 0.016)."),
       ("•  Drift removal falls with it:", " from 45.9 % to 0.2 % on cohort A, and from 62.5 % to 11.1 % on cohort B."),
       ("•  Cohort C:", " slowing costs 0.016 over two replication seeds (p = 0.003)."),
       ("•  Cohort B:", " slowing recovers 0.118, but the drift it still removes does not turn into accuracy.")]
write(get(s, "TextBox 4"), [P(R(b, 19, bold=True), R(t, 19, color=GRY), sa=13, ls=1.18) for b, t in bul])
write(get(s, "Rectangle 5"), [P(R("CONDITION   ", 14, bold=True),
                                R("The estimator's memory must outlast a class block. How much longer it can be before "
                                  "drift tracking is lost was not measured on these cohorts.", 18, italic=True),
                                sa=8, ls=1.15)])

# ------------------------------------------------------------------ 19 ablation
s = clone(9, skip=("Rectangle 6", "Rectangle 7", "Rectangle 8"))
title(s, "Component Ablation")
write(get(s, "TextBox 4"), [P(R("Cohort A, three data-split seeds under the current estimator. Rows marked pending await "
                                "seeds that are still running.", 13.5, italic=True, color=GRY))])
picture(s, ASSETS / "tab_ablation3.png", 1.6, 1.95, 10.1, 3.15)
note = get(s, "TextBox 9")
place(note, top=5.12)
write(note, [P(R("FA/min: spurious activations per minute of standing. Acc@90: accuracy at 90 % coverage.",
                 11.5, italic=True, color=GRY))])
tk = get(s, "TextBox 10")
place(tk, top=5.5, height=1.6)
write(tk, [lead("Alignment:  ", "+0.026 over the gate-only arm, higher in 6 of 7 participants (p = 0.22), a small gain "
                               "where class blocks are short.", 16),
           lead("Selective head and context:  ", "the head adds nothing (0.869 without, 0.868 with); context cuts false "
                                                  "activations by 27 % at a cost in calibration.", 16)])

# ------------------------------------------------------------------ 20 negative results
s = clone(3)
title(s, "Results Reported as Negative")
for name, num, head, body in (
        ("Rectangle 4", "1", "The selective head",
         "Declining the least-confident 10 % lowers accuracy in 14 of 18 arms, and the head adds no accuracy over three seeds."),
        ("Rectangle 5", "2", "Parameter efficiency",
         "EEG-TCNet reaches 0.878 with 7,062 parameters, against 0.868 with 24,181 for the proposed decoder."),
        ("Rectangle 6", "3", "The drift ceiling",
         "Two estimates of the drift timescale failed to resolve it, so only the lower bound of the rate condition is claimed.")):
    write(get(s, name), [P(R(num, 30, bold=True), sa=6, ls=1.1), P(R(head, 20, bold=True), sa=10, ls=1.05),
                         P(R(body, 16, color=GRY), sa=8, ls=1.22)])

# ------------------------------------------------------------------ 21 limitations
s = clone(7)
title(s, "Limitations and Next Steps")
bul = [("•  Small cohorts:", " seven, eight and twenty participants, so most paired tests have limited power."),
       ("•  Upper bound untested:", " no cohort samples drift finely enough to measure the drift timescale."),
       ("•  One participant per scalp map:", " the topography shows spatial separation; the ICA control carries the artefact test."),
       ("•  Baselines untuned:", " published decoders used our settings, and three failed to train on cohort C.")]
write(get(s, "TextBox 4"), [P(R(b, 19, bold=True), R(t, 19, color=GRY), sa=13, ls=1.18) for b, t in bul])
write(get(s, "Rectangle 5"), [P(R("NEXT   ", 14, bold=True),
                                R("Pre-register a protocol with class blocks of 100 to 200 windows and closely spaced "
                                  "recordings, where both bounds of the rate condition can be tested.", 18, italic=True),
                                sa=8, ls=1.15)])

# ------------------------------------------------------------------ 22 conclusion
s = clone(2)
title(s, "Conclusion")
write(get(s, "TextBox 4"), [P(R("A label-free alignment layer inside an EEG decoder removes session drift and keeps accuracy "
                                "level with published decoders. A running estimate, however, tracks whatever is being "
                                "streamed.", 20), sa=8, ls=1.25)])
write(get(s, "Rectangle 5"), [P(R("Its memory must outlast a class block, a rule read from the recording protocol before "
                                  "any model is trained.", 22, bold=True), sa=8, ls=1.15)])
write(get(s, "TextBox 6"), [P(R("The rule predicted a new cohort in advance.  ", 19, bold=True),
                              R("Where it holds, alignment helps a little; where it fails, alignment costs 0.149.",
                                19, color=GRY), sa=8, ls=1.2)])

# ------------------------------------------------------------------ 23 references
s = clone(11)
write(get(s, "TextBox 4"), [P(R("Works cited on these slides, keyed to their [author, year] tags.", 12.5, italic=True, color=GRY))])
left = [
    ("[Sarkar 2026]", "S. Sarkar et al., “EEG-controlled exoskeleton for walking and standing: a longitudinal multimodal dataset,” Sci. Data, 2026. (OpenNeuro ds007788)"),
    ("[Ferrero 2024]", "L. Ferrero et al., “Brain-machine interface based on deep learning to control a lower-limb robotic exoskeleton,” J. NeuroEng. Rehabil., 21(48), 2024."),
    ("[Soriano 2025]", "P. Soriano-Segura, M. Ortiz et al., “Characterization of error-related potentials during exoskeleton command via deep learning,” J. NeuroEng. Rehabil., 2025."),
    ("[Tortora 2023]", "S. Tortora et al., “Cortical and muscular activity under robot-assisted gait modes during exoskeleton walking,” 2023."),
    ("[Luu 2017]", "T. P. Luu et al., “Real-time EEG-based brain-computer interface to a virtual avatar enhances cortical involvement in human treadmill walking,” Sci. Rep., 7:8895, 2017."),
    ("[Schalk 2004]", "G. Schalk et al., “BCI2000: a general-purpose brain-computer interface (BCI) system,” IEEE TBME, 51(6):1034-1043, 2004."),
    ("[Lawhern 2018]", "V. J. Lawhern et al., “EEGNet: a compact CNN for EEG-based brain-computer interfaces,” J. Neural Eng., 15(5), 2018."),
    ("[Schirrmeister 2017]", "R. T. Schirrmeister et al., “Deep learning with convolutional neural networks for EEG decoding and visualization,” Hum. Brain Mapp., 38(11):5391-5420, 2017."),
]
right = [
    ("[Song 2023]", "Y. Song et al., “EEG Conformer: convolutional transformer for EEG decoding and visualization,” IEEE TNSRE, 31:710-719, 2023."),
    ("[He 2020]", "H. He and D. Wu, “Transfer learning for brain-computer interfaces: a Euclidean space data alignment approach,” IEEE TBME, 67(2):399-410, 2020."),
    ("[Rodrigues 2019]", "P. L. C. Rodrigues et al., “Riemannian Procrustes analysis: transfer learning for brain-computer interfaces,” IEEE TBME, 66(8):2390-2401, 2019."),
    ("[Barachant 2012]", "A. Barachant et al., “Multiclass brain-computer interface classification by Riemannian geometry,” IEEE TBME, 59(4), 2012."),
    ("[Wang 2021]", "D. Wang et al., “Tent: fully test-time adaptation by entropy minimization,” ICLR, 2021."),
    ("[Gong 2022]", "T. Gong et al., “NOTE: robust continual test-time adaptation against temporal correlation,” NeurIPS, 2022."),
    ("[Geifman 2019]", "Y. Geifman and R. El-Yaniv, “SelectiveNet: a deep neural network with an integrated reject option,” ICML, 2019."),
    ("[Pion-Tonachini 2019]", "L. Pion-Tonachini et al., “ICLabel: an automated electroencephalographic independent component classifier, dataset, and website,” NeuroImage, 198:181-197, 2019."),
]
for name, refs in (("TextBox 5", left), ("TextBox 6", right)):
    write(get(s, name), [P(R(tag + "  ", 11, bold=True), R(txt, 10.5, color=GRY), sa=9, ls=1.06) for tag, txt in refs])

# ------------------------------------------------------------------ drop the original proposal slides
ids = prs.slides._sldIdLst
for sid in list(ids)[:N_ORIG]:
    prs.part.drop_rel(sid.rId)
    ids.remove(sid)
prs.save(str(OUT))
print("saved %s | %d slides" % (OUT.name, len(prs.slides)))
