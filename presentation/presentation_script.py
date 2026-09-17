"""Slide-by-slide speaking script for CALM-Net_midterm_presentation_final.pptx.

Writes presentation_script.docx as a two-column table (slide number, script)
and presentation_script.md; the deck itself carries no speaker notes.

    python presentation/presentation_script.py
"""
from pathlib import Path

HERE = Path(__file__).resolve().parent
DECK = HERE / "CALM-Net_midterm_presentation_final.pptx"

SCRIPT = [
    ("Title", 20, [
        "Good morning. We are Md Wahiduzzaman Suva and Esme Moula Chowdhury Abha, and our supervisor is "
        "Dr. Muhammad Hasibur Rashid Chayon.",
        "Our project is CALM-Net. Today we present our midterm progress on label-free covariance alignment "
        "for longitudinal EEG decoding.",
    ]),
    ("Introduction", 40, [
        "Lower-limb exoskeletons controlled by EEG can help people with spinal cord injury or stroke walk again. "
        "EEG is non-invasive, so it requires no surgery.",
        "In practice, a decoder is trained once but must keep working across sessions recorded over several weeks. "
        "Between sessions, electrode placement and impedance change, and this alters the statistics of the EEG "
        "signal that the decoder was trained on.",
        "In this work, we propose a compact decoder that adapts to each new session without labels, and we "
        "evaluate it on three EEG datasets.",
    ]),
    ("Research Challenges", 40, [
        "We address three challenges.",
        "The first is session drift: the signal statistics shift from one session to the next.",
        "The second is movement artefact: walking produces motion and muscle activity that can appear in the EEG "
        "recording.",
        "The third is reliable decisions: exoskeleton control needs consistent accuracy across sessions and "
        "well-calibrated confidence.",
    ]),
    ("Research Objectives", 40, [
        "Our first objective is an adaptive alignment layer that corrects session drift without labels, both "
        "during training and at test time.",
        "The second is accurate and calibrated decoding, with accuracy comparable to published decoders from a "
        "compact model.",
        "The third is a design rule, read from the recording protocol, that sets how fast the alignment layer "
        "should adapt.",
        "Together these answer our research question: can a compact decoder adapt to session drift without labels "
        "while remaining accurate and reliable?",
    ]),
    ("Literature Review", 55, [
        "We reviewed related work in four areas. The first two are exoskeleton BCI and EEG decoding models.",
        "Sarkar and colleagues released the NeuroRex dataset, which we use as cohort A: seven healthy participants "
        "walking and standing with an EEG-controlled exoskeleton over nine sessions, with synchronised EEG, EOG, "
        "inertial and exoskeleton recordings.",
        "Ferrero and colleagues achieved closed-loop, asynchronous walk and stop control of a lower-limb exoskeleton "
        "with a deep-learning decoder, and used transfer learning to shorten the calibration in each session.",
        "Soriano-Segura and colleagues characterised error-related potentials during exoskeleton commands and "
        "detected them with deep learning, as a safety signal. Tortora and colleagues examined how cortical and "
        "muscular activity change across robot-assisted gait modes.",
        "For decoding, Schirrmeister and colleagues showed that convolutional networks decode movement-related EEG, "
        "with the shallow network learning band-power features similar to filter-bank common spatial patterns. "
        "EEGNet uses depthwise and separable convolutions to reduce the parameter count by an order of magnitude, "
        "and EEG Conformer adds a transformer encoder to a convolutional stem.",
        "In summary, exoskeleton control from EEG is well established, and compact convolutional decoders are the "
        "standard models. Their spatial filters are learned during training and stay fixed in later sessions.",
    ]),
    ("Literature Review (Continued)", 55, [
        "The other two areas are covariance-based alignment and test-time adaptation.",
        "Barachant and colleagues represented each EEG trial by its covariance matrix and classified trials by "
        "Riemannian distance to the class means.",
        "He and Wu proposed Euclidean alignment, which whitens the trials of each subject or session by the inverse "
        "square root of their mean covariance, so recordings become comparable without labels. Riemannian "
        "Procrustes analysis, by Rodrigues and colleagues, re-centres, stretches and rotates covariance "
        "distributions to transfer a classifier between subjects and sessions.",
        "In test-time adaptation, Tent updates normalisation parameters on unlabelled test data by minimising "
        "prediction entropy, and NOTE handles temporally correlated test streams with instance-aware batch "
        "normalisation and a prediction-balanced memory. SelectiveNet trains a classifier together with a selection "
        "head that abstains on uncertain inputs.",
        "The research gap is that existing alignment methods operate offline on complete recordings. Our work "
        "places alignment inside the network, adapts it online without labels, and provides a rule for choosing "
        "its adaptation rate.",
    ]),
    ("Proposed Architecture", 60, [
        "This is our proposed decoder. The input is 60 channels of EEG in four-second windows at 100 Hz.",
        "The first block is adaptive alignment. The layer keeps a running estimate of the covariance matrix M and "
        "whitens each window by M to the power of minus one half. A learned gate g combines the aligned and the "
        "original signal. The estimate updates without labels, including during inference.",
        "The second block is a multi-scale power stem. Temporal convolutions at 64, 128 and 256 milliseconds are "
        "followed by depthwise spatial filters, then squaring, pooling and a logarithm, which produces band-power "
        "features.",
        "An optional context module, a transformer over eight consecutive windows, can use longer temporal context.",
        "Finally, the selective head outputs the walk or stop decision with its confidence.",
        "The default configuration has 24,181 parameters.",
    ]),
    ("Adaptive Alignment Method", 50, [
        "The alignment layer maintains an exponential moving average of the spatial covariance of incoming windows.",
        "Each window is multiplied by the inverse square root of that estimate before feature extraction, which "
        "removes session-level differences in the covariance.",
        "Because the update needs no labels, the estimate keeps adapting when a new session begins.",
        "The adaptation rate sets how many windows the estimate remembers. Our rule is that this memory should be "
        "longer than the length of a single-class block in the recording protocol, so the estimate reflects the "
        "session and not the class being performed.",
        "The key idea is that session drift is corrected inside the network, and the recording protocol determines "
        "the adaptation rate before training.",
    ]),
    ("Datasets and Experimental Setup", 45, [
        "We use three datasets.",
        "Cohort A is an exoskeleton dataset with 7 participants and 9 sessions over several weeks. We train on "
        "sessions 1 to 3 and test on sessions 4 to 9.",
        "Cohort B is a treadmill walking dataset with 8 participants, recorded in an independent laboratory.",
        "Cohort C is the PhysioNet motor execution dataset with 20 participants. For this cohort, we registered our "
        "analysis plan before training any model.",
        "All models use the same protocol: 8 to 30 Hz filtering, AdamW optimisation with early stopping, and three "
        "data-split seeds. Seven published decoders were trained with identical settings for comparison.",
    ]),
    ("Session Drift Reduction", 40, [
        "First, we measured how much session drift the alignment layer removes. The bars show the Riemannian "
        "distance between the training-session covariance and each test session.",
        "Orange is the raw signal, grey is a control that never updates its estimate, and black is the aligned "
        "signal.",
        "Alignment reduces the distance by 45.9 percent on cohort A and 62.5 percent on cohort B, and the reduction "
        "appears in all 15 participants. The control leaves the distance unchanged, which shows that the reduction "
        "comes from adaptation.",
    ]),
    ("Per-Participant Drift Reduction", 25, [
        "This table gives the values for every participant.",
        "On cohort A, the reduction ranges from 38.7 to 51.2 percent, and on cohort B from 43.8 to 82.1 percent.",
    ]),
    ("Comparison with Baseline Models", 60, [
        "Here we compare our decoder with published models, all trained in the same pipeline and averaged over "
        "three data-split seeds.",
        "On cohort A, our decoder reaches 0.868 balanced accuracy, comparable to the best baseline at 0.876, and it "
        "uses fewer parameters than six of the seven baselines.",
        "On cohort C, our decoder has the highest accuracy, 0.796.",
        "On cohort B, the protocol rule selects the configuration without alignment, which reaches 0.729, second of "
        "the eight models.",
        "Calibration error also stays low, 0.058 on cohort A and 0.061 on cohort C.",
    ]),
    ("Alignment Estimate During Recording", 45, [
        "This figure shows why the adaptation rate matters. The top row traces the alignment estimate during a "
        "recording, and the bottom row shows its distance to the walk and stop covariances.",
        "On cohort A, where class blocks are short, the estimate stays between the two classes.",
        "On cohort B, where class blocks are long, a fast estimate follows the class currently being performed. "
        "This is exactly the situation our adaptation rate rule is designed to detect.",
    ]),
    ("Adaptation Rate Analysis", 40, [
        "We tested the rule on cohort B by changing the adaptation rate.",
        "As the estimator memory increases from 160 to 3200 windows, accuracy rises from 0.580 to 0.698, as the "
        "rule predicts.",
        "Because cohort B has class blocks longer than the default memory, the rule selects the configuration "
        "without alignment, which reaches 0.729.",
    ]),
    ("Prospective Validation on Cohort C", 40, [
        "To test the rule in advance, we registered two predictions for cohort C before training any model on it.",
        "The first prediction was that enabling alignment would not reduce accuracy. The second was that a slower "
        "adaptation rate would bring no change, because cohort C has short class blocks.",
        "Both predictions were confirmed: accuracy changed by plus 0.002 with alignment, and the slower rate made "
        "no meaningful difference. Two additional data-split seeds support the same conclusions.",
    ]),
    ("Ablation Study", 35, [
        "The ablation study on cohort A shows the contribution of each component, averaged over three seeds.",
        "Adding the alignment layer raises balanced accuracy from 0.842 with the gate alone to 0.868.",
        "The cross-epoch context module reduces spurious activations from 1.60 to 1.16 per minute of standing, "
        "which is valuable for safe exoskeleton control.",
    ]),
    ("Artefact Analysis", 40, [
        "Because walking can introduce artefacts, we retrained the decoder after removing non-neural components "
        "with independent component analysis and ICLabel.",
        "After cleaning, mean balanced accuracy is 0.785, well above chance, and the change for each participant "
        "does not follow the number of muscle components removed.",
        "These results are consistent with the decoder using neural EEG activity.",
    ]),
    ("Scalp Topography", 30, [
        "This topography shows one participant. The right panel compares 8 to 30 Hz power during walking and "
        "standing.",
        "Power decreases over sensorimotor electrodes during walking, which is consistent with mu and beta "
        "desynchronisation during movement.",
    ]),
    ("Limitations and Future Work", 35, [
        "Our current work has some limitations.",
        "The datasets include 7, 8 and 20 participants. All results come from recorded datasets, and closed-loop "
        "testing is planned. The architecture is still under development, and we will evaluate further variants. "
        "The adaptation rate is currently set by a rule and could be learned from data.",
        "In future work, we plan real-time evaluation with an exoskeleton, larger cohorts, and recording protocols "
        "designed to refine the adaptation rate rule.",
    ]),
    ("Conclusion", 30, [
        "To conclude, we presented a compact EEG decoder with an adaptive alignment layer that corrects session "
        "drift without labels.",
        "The decoder reduces session drift by 46 to 63 percent, achieves accuracy comparable to published decoders, "
        "and follows a protocol-based rule that was confirmed prospectively.",
        "This approach supports longitudinal use of EEG-controlled exoskeletons with less recalibration.",
    ]),
    ("References", 10, [
        "These are the works we cited. Thank you for your attention. We are happy to take your questions.",
    ]),
]


def total_seconds():
    return sum(t for _, t, _ in SCRIPT)


def write_md():
    lines = ["# CALM-Net midterm presentation: speaking script", "",
             "Deck: `CALM-Net_midterm_presentation_final.pptx`, %d slides, about %d minutes in total."
             % (len(SCRIPT), round(total_seconds() / 60)), ""]
    for i, (name, secs, paras) in enumerate(SCRIPT, 1):
        lines += ["## Slide %d. %s  (about %d s)" % (i, name, secs), ""]
        lines += [p + "\n" for p in paras]
    (HERE / "presentation_script.md").write_text("\n".join(lines), encoding="utf-8")


def write_docx():
    """Two-column table: slide number, script."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    doc.add_heading("CALM-Net Midterm Presentation Script", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ("Slide", "Script")):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
    for i, (name, secs, paras) in enumerate(SCRIPT, 1):
        row = table.add_row().cells
        row[0].text = str(i)
        row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row[1].text = ""
        for k, para in enumerate(paras):
            p = row[1].paragraphs[0] if k == 0 else row[1].add_paragraph()
            p.add_run(para)
            p.paragraph_format.space_after = Pt(4)
    for row in table.rows:
        row.cells[0].width = Cm(2.0)
        row.cells[1].width = Cm(14.5)
    doc.save(str(HERE / "presentation_script.docx"))


def clear_notes(path=DECK):
    """Remove any speaker notes from the deck."""
    from pptx import Presentation
    prs = Presentation(str(path))
    for slide in prs.slides:
        if slide.has_notes_slide:
            slide.notes_slide.notes_text_frame.text = ""
    prs.save(str(path))
    return path.name


def write_notes(path=DECK):
    from pptx import Presentation
    prs = Presentation(str(path))
    assert len(prs.slides) == len(SCRIPT), (len(prs.slides), len(SCRIPT))
    for slide, (name, secs, paras) in zip(prs.slides, SCRIPT):
        slide.notes_slide.notes_text_frame.text = "\n\n".join(paras)
    try:
        prs.save(str(path))
        return path.name
    except PermissionError:
        alt = path.with_name(path.stem + "_with_notes.pptx")
        prs.save(str(alt))
        return alt.name


if __name__ == "__main__":
    write_md()
    write_docx()
    print("script: %d slides, %.1f minutes" % (len(SCRIPT), total_seconds() / 60))
    print("speaker notes cleared in", clear_notes())
